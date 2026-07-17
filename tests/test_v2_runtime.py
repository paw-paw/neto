from __future__ import annotations

import re
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

from parser import parse_workbook
from parser.models import ParsedMatch, ParserKey
from parser.v2_runtime import _add_duplicates


def _pipeline(
    source: dict,
    *,
    on_missing: str = "blocking_error",
    default=None,
    fallbacks: list[dict] | None = None,
) -> dict:
    return {
        "value": {
            "source": source,
            "transforms": [],
            "fallbacks": fallbacks or [],
        },
        "on_missing": on_missing,
        "default": default,
    }


def _literal(value) -> dict:
    return {"op": "literal.value", "args": {"value": value}}


def _cell(column: str) -> dict:
    return {"op": "cell.column", "args": {"column": column}}


def _v2_key(
    *,
    team_a_on_missing: str = "use_default",
    team_a_default: str | None = "TBD",
    team_a_fallbacks: list[dict] | None = None,
    time_source: dict | None = None,
    time_on_missing: str = "blocking_error",
    count_minimum: int | None = 1,
    count_maximum: int | None = 1,
    count_severity: str = "warning",
    formula_severity: str = "warning",
    duplicate_config: dict | None = None,
) -> ParserKey:
    raw_data = {
        "metadata": {"source_files": [{"filename": "runtime-test.xlsx"}]},
        "workbook": {
            "formula_value_policy": {
                "mode": "first_available",
                "on_missing_cached_value": formula_severity,
            },
            "hidden_content_policy": {"include_hidden_rows": False},
        },
        "sources": [
            {
                "source_id": "main",
                "sheet_locator": {
                    "op": "sheet.exact",
                    "args": {"sheet_name": "Schedule"},
                },
                "required": True,
                "value_mode": "inherit",
            }
        ],
        "record_sets": [
            {
                "record_set_id": "matches",
                "enabled": True,
                "source_id": "main",
                "locator": {
                    "op": "records.row_ranges",
                    "args": {
                        "anchor_column": "A",
                        "ranges": [{"start_row": 2, "end_row": 2}],
                    },
                },
                "fields": {
                    "date": _pipeline(_literal("2026-07-01")),
                    "time": _pipeline(
                        time_source or _literal("18:00"),
                        on_missing=time_on_missing,
                    ),
                    "team_a": _pipeline(
                        _cell("A"),
                        on_missing=team_a_on_missing,
                        default=team_a_default,
                        fallbacks=team_a_fallbacks,
                    ),
                    "team_b": _pipeline(_cell("B")),
                    "stage": _pipeline(_literal("Group Stage"), on_missing="null"),
                    "bo": _pipeline(_literal("BO3"), on_missing="null"),
                    "match_label": _pipeline(_literal("M1"), on_missing="null"),
                },
            }
        ],
        "normalization": {
            "datetime": {"ambiguous_time_policy": "blocking_error"},
            "teams": {
                "placeholder_values": ["TBD", "TBA"],
                "placeholder_patterns": [r"^Winner M\d+$"],
            },
            "best_of": {"normalize": True},
        },
        "validation": {
            "record_count": {
                "minimum": count_minimum,
                "maximum": count_maximum,
                "on_violation": count_severity,
            },
            "duplicate_check": duplicate_config
            or {
                "enabled": False,
                "severity": "warning",
                "fields": ["start_time_utc", "team_a", "team_b"],
                "team_order_sensitive": True,
            },
        },
    }
    return ParserKey(
        parser_key_id="runtime_test_v2",
        key_name="Runtime Test v2",
        tournament_name="Runtime Test",
        base_timezone="America/Lima",
        target_sheet="Schedule",
        layout_type="operator_graph",
        header_row=1,
        data_start_row=2,
        field_mappings={},
        forward_fill_rules={},
        schema_version="neto.parser_key.v2",
        raw_data=raw_data,
    )


def _workbook_bytes(sheets: dict[str, dict[str, object]]) -> bytes:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for sheet_name, cells in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for coordinate, value in cells.items():
            sheet[coordinate] = value
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _with_cached_formula(data: bytes, coordinate: str, value: str) -> bytes:
    source = BytesIO(data)
    output = BytesIO()
    with ZipFile(source) as input_zip, ZipFile(
        output, "w", compression=ZIP_DEFLATED
    ) as output_zip:
        for info in input_zip.infolist():
            payload = input_zip.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                pattern = rb'<c r="' + coordinate.encode() + rb'"[^>]*>.*?</c>'
                replacement = (
                    b'<c r="'
                    + coordinate.encode()
                    + b'" t="str"><f>C2</f><v>'
                    + value.encode()
                    + b"</v></c>"
                )
                payload, count = re.subn(pattern, replacement, payload, count=1)
                assert count == 1
            output_zip.writestr(info, payload)
    return output.getvalue()


def _assert_no_internal_values(result) -> None:
    for match in result.matches:
        for value in match.as_output_dict().values():
            assert not str(value).lstrip().startswith("=")
            assert "openpyxl" not in str(value).casefold()


def test_cached_formula_uses_cached_value_without_missing_cache_warning() -> None:
    data = _workbook_bytes(
        {"Schedule": {"A2": "=C2", "B2": "Opponent", "C2": "Target"}}
    )
    data = _with_cached_formula(data, "A2", "Cached Team")

    result = parse_workbook(data, _v2_key())

    assert result.status == "parsed"
    assert result.matches[0].team_a == "Cached Team"
    assert "formula_cached_value_missing" not in {
        issue.code for issue in result.issues
    }


def test_direct_local_reference_and_reference_chain_are_resolved() -> None:
    result = parse_workbook(
        _workbook_bytes(
            {
                "Schedule": {
                    "A2": "=$C$2",
                    "B2": "Opponent",
                    "C2": "=D2",
                    "D2": "Resolved Team",
                }
            }
        ),
        _v2_key(),
    )

    assert result.status == "parsed_with_warnings"
    assert result.matches[0].team_a == "Resolved Team"
    provenance = result.matches[0].field_provenance["team_a"]
    assert provenance["formula_text"] == "=$C$2"
    assert provenance["formula_resolution"] == "resolved_reference"
    assert provenance["resolved_source_cell"] == "Schedule!D2"


def test_direct_cross_sheet_reference_is_resolved() -> None:
    result = parse_workbook(
        _workbook_bytes(
            {
                "Schedule": {"A2": "='Teams Data'!$C$2", "B2": "Opponent"},
                "Teams Data": {"C2": "Cross-sheet Team"},
            }
        ),
        _v2_key(),
    )

    assert result.matches[0].team_a == "Cross-sheet Team"
    assert (
        result.matches[0].field_provenance["team_a"]["resolved_source_cell"]
        == "Teams Data!C2"
    )


@pytest.mark.parametrize(
    "formula",
    ["=[Other.xlsx]Sheet1!A1", "=UPPER(C2)", "=C2:D2", "=NamedTeam"],
)
def test_unsupported_formulas_use_explicit_tbd_policy(formula: str) -> None:
    result = parse_workbook(
        _workbook_bytes(
            {"Schedule": {"A2": formula, "B2": "Opponent", "C2": "Target"}}
        ),
        _v2_key(),
    )

    assert result.status == "parsed_with_warnings"
    assert result.matches[0].team_a == "TBD"
    assert "formula_cached_value_missing" in {issue.code for issue in result.issues}
    assert result.matches[0].field_provenance["team_a"]["used_default"] is True
    _assert_no_internal_values(result)


def test_empty_reference_and_cycle_use_tbd_with_formula_provenance() -> None:
    empty = parse_workbook(
        _workbook_bytes(
            {"Schedule": {"A2": "=C2", "B2": "Opponent", "C2": None}}
        ),
        _v2_key(),
    )
    cycle = parse_workbook(
        _workbook_bytes(
            {"Schedule": {"A2": "=C2", "B2": "Opponent", "C2": "=A2"}}
        ),
        _v2_key(),
    )

    assert empty.matches[0].team_a == cycle.matches[0].team_a == "TBD"
    assert (
        empty.matches[0].field_provenance["team_a"]["formula_resolution"]
        == "referenced_cell_empty"
    )
    assert (
        cycle.matches[0].field_provenance["team_a"]["formula_resolution"]
        == "cycle_detected"
    )


def test_array_formula_object_never_reaches_canonical_output() -> None:
    formula = ArrayFormula(ref="A2:A2", text="=C2")
    result = parse_workbook(
        _workbook_bytes(
            {
                "Schedule": {
                    "A2": formula,
                    "B2": "Opponent",
                    "C2": "Should not resolve",
                }
            }
        ),
        _v2_key(),
    )

    assert result.matches[0].team_a == "TBD"
    assert (
        result.matches[0].field_provenance["team_a"]["formula_resolution"]
        == "array_formula_not_supported"
    )
    _assert_no_internal_values(result)


def test_blocking_field_remains_blocking_when_formula_cannot_be_resolved() -> None:
    result = parse_workbook(
        _workbook_bytes(
            {"Schedule": {"A2": "=UPPER(C2)", "B2": "Opponent", "C2": "Team"}}
        ),
        _v2_key(team_a_on_missing="blocking_error", team_a_default=None),
    )

    assert result.status == "blocked"
    assert result.matches[0].team_a == ""
    assert {"missing_team_a", "formula_cached_value_missing"} <= {
        issue.code for issue in result.issues
    }
    _assert_no_internal_values(result)


def test_formula_diagnostic_survives_pipeline_fallback() -> None:
    result = parse_workbook(
        _workbook_bytes(
            {"Schedule": {"A2": "=UPPER(C2)", "B2": "Opponent", "C2": "Team"}}
        ),
        _v2_key(
            team_a_on_missing="blocking_error",
            team_a_default=None,
            team_a_fallbacks=[
                {
                    "source": _literal("TBD"),
                    "transforms": [],
                    "fallbacks": [],
                }
            ],
        ),
    )

    assert result.status == "parsed_with_warnings"
    assert result.matches[0].team_a == "TBD"
    provenance = result.matches[0].field_provenance["team_a"]
    assert provenance["used_fallback"] is True
    assert provenance["fallback_reason"] == "formula_cache_missing"
    assert "formula_cached_value_missing" in {issue.code for issue in result.issues}


def test_tbd_time_is_exportable_as_time_pending() -> None:
    result = parse_workbook(
        _workbook_bytes({"Schedule": {"A2": "Team", "B2": "Opponent", "C2": "TBD"}}),
        _v2_key(time_source=_cell("C"), time_on_missing="blocking_error"),
    )

    assert result.status == "parsed_with_warnings"
    assert result.matches[0].time_original == ""
    assert result.matches[0].start_time_utc == ""
    assert [issue.severity for issue in result.issues if issue.code == "time_pending"] == [
        "warning"
    ]


@pytest.mark.parametrize(
    ("minimum", "maximum", "severity", "expected_status"),
    [
        (2, None, "warning", "parsed_with_warnings"),
        (2, None, "blocking_error", "blocked"),
        (None, 0, "warning", "parsed_with_warnings"),
        (None, None, "blocking_error", "parsed"),
    ],
)
def test_record_count_honors_severity_and_null_limits(
    minimum, maximum, severity, expected_status
) -> None:
    result = parse_workbook(
        _workbook_bytes({"Schedule": {"A2": "Team", "B2": "Opponent"}}),
        _v2_key(
            count_minimum=minimum,
            count_maximum=maximum,
            count_severity=severity,
        ),
    )

    assert result.status == expected_status
    assert ("record_count_mismatch" in {issue.code for issue in result.issues}) is (
        minimum is not None or maximum is not None
    )


def _match(team_a: str, team_b: str, *, row: int = 2, time: str = "2026-07-01T00:00:00Z") -> ParsedMatch:
    return ParsedMatch(
        source_row=row,
        source_sheet="Schedule",
        record_set_id="matches",
        date_original="2026-07-01",
        time_original="00:00",
        timezone="UTC",
        start_time_utc=time,
        team_a=team_a,
        team_b=team_b,
        stage="",
        bo="",
        match_label="",
    )


def test_duplicate_rules_honor_placeholders_order_fields_and_severity() -> None:
    team_config = {
        "placeholder_values": ["TBD", "TBA"],
        "placeholder_patterns": [r"^Winner M\d+$"],
    }
    config = {
        "fields": ["start_time_utc", "team_a", "team_b"],
        "severity": "blocking_error",
        "team_order_sensitive": False,
    }
    issues = []
    _add_duplicates(
        [_match("Alpha", "TBD", row=2), _match("TBD", "Alpha", row=3)],
        issues,
        config,
        team_config,
    )
    assert len(issues) == 2
    assert {issue.severity for issue in issues} == {"blocking_error"}

    issues = []
    _add_duplicates(
        [_match("Alpha", "TBD", row=2), _match("TBD", "Alpha", row=3)],
        issues,
        {**config, "team_order_sensitive": True},
        team_config,
    )
    assert issues == []

    issues = []
    _add_duplicates(
        [_match("TBD", "TBA", row=2), _match("TBA", "TBD", row=3)],
        issues,
        config,
        team_config,
    )
    assert issues == []

    issues = []
    _add_duplicates(
        [_match("Alpha", "One", row=2), _match("Alpha", "Two", row=3)],
        issues,
        {
            "fields": ["start_time_utc", "team_a"],
            "severity": "warning",
            "team_order_sensitive": True,
        },
        team_config,
    )
    assert len(issues) == 2
    assert {issue.affected_field for issue in issues} == {"start_time_utc,team_a"}
